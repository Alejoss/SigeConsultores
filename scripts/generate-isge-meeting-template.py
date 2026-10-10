#!/usr/bin/env python3
"""Genera la plantilla Word importable para las reuniones de ISGE 360."""
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "client" / "public" / "plantillas" / "acta-reunion-isge.docx"
BLUE = "1F4E78"
LIGHT_BLUE = "DDEBF7"
LIGHT_GRAY = "F2F2F2"


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def label_row(table, label, placeholder, wide=False):
    row = table.add_row()
    if wide:
        label_cell = row.cells[0].merge(row.cells[1])
        value_cell = row.cells[2].merge(row.cells[3])
    else:
        label_cell = row.cells[0]
        value_cell = row.cells[1].merge(row.cells[3])
    shade(label_cell, LIGHT_BLUE)
    shade(value_cell, LIGHT_GRAY)
    cell_text(label_cell, label, bold=True, color=BLUE)
    cell_text(value_cell, placeholder, color="666666")
    return row


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("PLANTILLA DE ACTA DE REUNIÓN ISGE")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Formato único para importar una reunión, sus acuerdos y responsabilidades en ISGE 360")
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(9)

    notice = document.add_paragraph()
    notice.style = document.styles["Normal"]
    notice_run = notice.add_run(
        "Instrucciones: complete los campos grises y la tabla de acuerdos. No cambie los títulos ni las columnas. "
        "Las filas vacías no se importarán. Antes de confirmar, ISGE mostrará una vista previa para aprobar, editar, "
        "asignar o rechazar cada acuerdo."
    )
    notice_run.font.size = Pt(9)

    document.add_paragraph().add_run("DATOS DE LA REUNIÓN").bold = True
    details = document.add_table(rows=0, cols=4)
    style_table(details)
    widths = [Cm(4.2), Cm(3.5), Cm(4.2), Cm(3.5)]
    for row in details.rows:
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
    label_row(details, "Nombre de la reunión", "[Ej.: Revisión semanal de indicadores]", wide=True)
    label_row(details, "Fecha", "[DD/MM/AAAA]", wide=False)
    label_row(details, "Hora", "[HH:MM – HH:MM]", wide=False)
    label_row(details, "Lugar o medio", "[Presencial, Zoom, Teams, Meet…]", wide=True)
    label_row(details, "Objetivo", "[Objetivo principal de la reunión]", wide=True)
    label_row(details, "Participantes", "[Nombres y cargos o lista de participantes]", wide=True)
    label_row(details, "Resumen, desarrollo o decisiones", "[Síntesis de lo tratado y decisiones relevantes]", wide=True)

    document.add_paragraph()
    document.add_paragraph().add_run("ACUERDOS Y ACTIVIDADES").bold = True
    instructions = document.add_paragraph("Registre un acuerdo por fila. El responsable puede quedar vacío si se definirá después en la vista previa de ISGE.")
    instructions.runs[0].italic = True
    instructions.runs[0].font.size = Pt(8.5)

    agreements = document.add_table(rows=1, cols=5)
    style_table(agreements)
    headers = ["Acción o acuerdo", "Responsable", "Proceso responsable", "Fecha límite", "Observaciones"]
    for index, text in enumerate(headers):
        shade(agreements.rows[0].cells[index], BLUE)
        cell_text(agreements.rows[0].cells[index], text, bold=True, color="FFFFFF", size=8.5)
    for _ in range(8):
        row = agreements.add_row()
        placeholders = [
            "[Describa la acción o compromiso]",
            "[Nombre, si se conoce]",
            "[Proceso ISGE, si aplica]",
            "[DD/MM/AAAA]",
            "[Opcional]",
        ]
        for index, placeholder in enumerate(placeholders):
            shade(row.cells[index], LIGHT_GRAY)
            cell_text(row.cells[index], placeholder, color="777777", size=8.5)

    document.add_paragraph()
    final_note = document.add_paragraph()
    final_note.add_run("Nota de control: ").bold = True
    final_note.add_run(
        "subir esta plantilla no envía correos ni crea vínculos automáticamente. "
        "ISGE solicita revisar cada acuerdo antes de convertirlo en una responsabilidad oficial."
    )
    final_note.runs[-1].font.size = Pt(8.5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("ISGE 360 · Plantilla importable de acta de reunión")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string("666666")

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
